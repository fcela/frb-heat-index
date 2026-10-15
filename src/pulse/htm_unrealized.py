"""Bank HTM unrealized losses (aggregate US banking sector).

Source: FDIC Quarterly Banking Profile (QBP) time-series spreadsheet
(https://www.fdic.gov/quarterly-banking-profile) — aggregate workbook derived
from the same Call Report source as raw FFIEC bulk, pre-aggregated by the FDIC.
Workbook URL is discovered live because naming varies by quarter (not hardcoded).

Frequency: quarterly. Unit: USD billions (positive = losses).
The workbook's "Balance Sheet" sheet, row
"Held to maturity on non-equity securities unrealized gain/loss", is reported
in $ millions as a gain/loss (losses negative); value = -raw/1000 so positive
values are losses.
"""

import io
import re
import subprocess
from urllib.parse import urljoin

import polars as pl

QBP_URL = "https://www.fdic.gov/quarterly-banking-profile"
UA = "Mozilla/5.0 (research; financial-vulnerability-monitor)"


def _get(url: str, timeout: int = 90) -> bytes:
    r = subprocess.run(
        ["curl", "-sL", "--max-time", str(timeout), "-A", UA, url],
        capture_output=True,
    )
    if r.returncode != 0 or len(r.stdout) < 100:
        raise RuntimeError(f"FDIC fetch failed for {url}: curl rc={r.returncode}")
    return r.stdout


def discover_workbook_url() -> str:
    """Find the current QBP time-series spreadsheet URL from the FDIC site."""
    landing = _get(QBP_URL).decode("utf-8", errors="replace")
    m = re.search(
        r'href="(/quarterly-banking-profile/quarterly-banking-profile-[0-9]q-\d{4})"',
        landing,
    )
    if not m:
        raise RuntimeError("FDIC: could not find latest QBP report page")
    report = _get(urljoin(QBP_URL, m.group(1))).decode("utf-8", errors="replace")
    m2 = re.search(r'href="([^"]*time-series-spreadsheets[^"]*\.xlsx?)"', report)
    if not m2:
        raise RuntimeError("FDIC: could not find time-series spreadsheet link")
    return urljoin(QBP_URL, m2.group(1))


def fetch_htm_unrealized(timeout: int = 120) -> pl.DataFrame:
    """Fetch bank HTM unrealized losses. Returns date (quarter-start)/value."""
    url = discover_workbook_url()
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(_get(url, timeout=timeout)), read_only=True, data_only=True
    )
    if "Balance Sheet" not in wb.sheetnames:
        raise RuntimeError("FDIC: 'Balance Sheet' sheet missing")
    ws = wb["Balance Sheet"]
    # Header row: quarter labels like "1984Q1" starting at column C.
    header_row, quarters = None, []
    for rn in range(1, 15):
        vals = [c.value for c in ws[rn]]
        if any(
            isinstance(v, str) and re.fullmatch(r"\d{4}Q[1-4]", v.strip())
            for v in vals
        ):
            header_row = rn
            quarters = [str(v).strip() for v in vals]
            break
    if header_row is None:
        raise RuntimeError("FDIC: quarter header row not found")
    label_idx = quarters.index(next(q for q in quarters if q))
    htm_row = None
    for rn in range(header_row, header_row + 60):
        v = ws[rn][1].value
        if v and isinstance(v, str) and re.search(
            r"Held to maturity.*unrealized", v, re.IGNORECASE
        ):
            htm_row = rn
            break
    if htm_row is None:
        raise RuntimeError("FDIC: HTM unrealized row not found")
    recs = []
    for ci, q in enumerate(quarters):
        if not re.fullmatch(r"\d{4}Q[1-4]", q or ""):
            continue
        raw = ws[htm_row][ci].value
        try:
            val = float(str(raw).replace(",", "").strip())
        except (TypeError, ValueError):
            continue  # 'N/A' early years
        yy, qq = int(q[:4]), int(q[5])
        date = f"{yy:04d}-{3 * (qq - 1) + 1:02d}-01"
        # Workbook amounts in $ millions, gain/loss (losses negative).
        # Positive output = losses, in USD billions.
        val = -val / 1000.0
        if not (-1e9 < val < 1e9) or val != val:  # NaN/Inf guard
            continue
        recs.append((date, val))
    if not recs:
        raise RuntimeError("FDIC: no usable HTM observations")
    df = pl.DataFrame({"date": [d for d, _ in recs], "value": [v for _, v in recs]})
    df = df.with_columns(
        pl.col("date").str.to_date(), pl.col("value").cast(pl.Float64)
    ).sort("date")
    return df.select("date", "value")
